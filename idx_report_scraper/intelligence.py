"""Build the "Account Intelligence" workbook (Profil & Keuangan + Kalkulator
Kesehatan) from PDFs already downloaded by downloader.run()/bulk.run_all().

Scope is deliberately the deterministic half of the account-intelligence
process: figures that live in the financial-statement PDF itself. The
qualitative half (corporate-action news, event calendars, executive names,
communication angles - see the "Account_Intelligence" template this was
built to reproduce) needs actual internet research per company and stays a
separate, manual/AI-assisted step - not something a parser can fabricate.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import financials

logger = logging.getLogger("idx_report_scraper")

# Prefer the annual (audited) report when present; otherwise the most
# recently closed quarter available.
_PERIODE_PREFERENCE = ("tahunan", "tw3", "tw2", "tw1")

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SUBTITLE_FONT = Font(name=FONT_NAME, italic=True, size=10, color="595959")
INPUT_FONT = Font(name=FONT_NAME, color="0000FF")
NORMAL_FONT = Font(name=FONT_NAME)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM_FMT = "#,##0;(#,##0);-"
PCT_FMT = "0.0%;(0.0%);-"
RATIO_FMT = "0.00x;(0.00x);-"
WRAP = Alignment(wrap_text=True, vertical="top")

# The 16-combination health matrix, copied verbatim from the reference
# "Account Intelligence" template (Matriks Kesehatan sheet) so the
# Kalkulator Kesehatan formulas this module writes resolve against it.
HEALTH_MATRIX = [
    ("B", "B", "Byk", "B", "Kuat", "Raksasa sehat, padat karya, laba jadi kas",
     "Prioritas tinggi. Full-service + employer branding (karyawan banyak)"),
    ("B", "B", "Byk", "K", "Sehat", "Sehat, laba tertahan di capex/dividen; cek likuiditas",
     "Mampu; jaga termin, jual dgn ROI jelas; employer branding"),
    ("B", "B", "Sdk", "B", "Kuat", "Mesin uang padat modal, kas melimpah",
     "Prioritas tinggi. Fokus IR/brand/ESG, bukan employer branding"),
    ("B", "B", "Sdk", "K", "Sehat", "Efisien & untung, tapi kas terpakai",
     "Mampu tapi selektif; pitch berbasis nilai"),
    ("B", "K", "Byk", "B", "Sedang", "Skala besar, margin tipis, padat karya, kas kuat",
     "Mampu belanja; jual efisiensi + employer branding"),
    ("B", "K", "Byk", "K", "Waspada", "Margin tipis + kas tipis + beban gaji besar",
     "Siklus jual sulit; hanya proyek ROI tinggi"),
    ("B", "K", "Sdk", "B", "Sedang", "Padat modal, margin lesu (harga/investasi) tapi kas melimpah",
     "Mampu meski margin turun; narasi/IR/transisi energi"),
    ("B", "K", "Sdk", "K", "Waspada", "Besar tapi margin & kas tertekan",
     "Sangat selektif, fokus ROI cepat"),
    ("K", "B", "Byk", "B", "Sehat", "Margin tinggi tapi skala kecil & padat karya, kas kuat",
     "Butuh brand building untuk tumbuh; mampu bayar"),
    ("K", "B", "Byk", "K", "Sedang", "Untung tapi beban SDM berat & kas tipis",
     "Hati-hati termin; proyek kecil dulu"),
    ("K", "B", "Sdk", "B", "Sehat", "Niche premium, ramping, kas kuat",
     "Mampu; jual growth marketing & brand"),
    ("K", "B", "Sdk", "K", "Sehat", "Untung tinggi tapi kas baru/terpakai",
     "Selektif, proyek terukur"),
    ("K", "K", "Byk", "B", "Waspada", "Kecil, untung tipis, overstaffed, tapi ada kas",
     "Kas beri ruang, kinerja lemah; jangan prioritaskan"),
    ("K", "K", "Byk", "K", "Rawan", "Kecil, rugi/tipis, padat karya, kas tipis",
     "Prospek lemah; lewati"),
    ("K", "K", "Sdk", "B", "Potensi", "Kecil & belum matang tapi ramping & kas kuat (baru IPO/didanai)",
     "Bisa tumbuh; kas ada meski kinerja muda"),
    ("K", "K", "Sdk", "K", "Rawan", "Kecil di segala hal",
     "Lemah kecuali early-stage menjanjikan"),
]


def _style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _style_data_row(ws, row, ncols, wrap_cols=()):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = NORMAL_FONT
        cell.border = BORDER
        if c in wrap_cols:
            cell.alignment = WRAP


def _title_block(ws, title, subtitle, ncols, subtitle_height=28):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    last_col = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    ws.row_dimensions[2].height = subtitle_height
    ws["A2"].alignment = WRAP


def discover_downloaded_tickers(out_dir: Path, year: int) -> list[str]:
    """List every ticker under out_dir/<year>/ that has a manifest.json -
    i.e. every company downloader.run()/bulk.run_all() has already
    produced results for, without needing a separate tickers file."""
    year_dir = Path(out_dir) / str(year)
    if not year_dir.is_dir():
        return []
    return sorted(
        p.name for p in year_dir.iterdir() if p.is_dir() and (p / "manifest.json").exists()
    )


def pick_best_pdf(out_dir: Path, year: int, kode_emiten: str) -> tuple[str | None, Path | None]:
    """Return (periode, pdf_path) for the most useful PDF a company has
    downloaded - the annual report if present, else the latest available
    quarter - or (None, None) if nothing was downloaded for it."""
    manifest_path = out_dir / str(year) / kode_emiten / "manifest.json"
    if not manifest_path.exists():
        return None, None

    entries = json.loads(manifest_path.read_text(encoding="utf-8"))
    by_periode = {e["periode"]: e for e in entries}
    for periode in _PERIODE_PREFERENCE:
        entry = by_periode.get(periode)
        if entry is None:
            continue
        pdf_path = Path(entry["local_path"])
        if pdf_path.exists():
            return periode, pdf_path
    return None, None


def collect_financial_profiles(
    out_dir: Path, year: int, kode_emiten_list: Iterable[str]
) -> list[tuple[str, str, financials.FinancialProfile]]:
    """Extract a FinancialProfile for every ticker in kode_emiten_list that
    has a usable downloaded PDF. Returns (kode_emiten, periode, profile)
    triples; tickers with nothing downloaded, or whose PDF fails to parse,
    are skipped with a log message rather than aborting the batch.
    """
    results = []
    for kode in kode_emiten_list:
        periode, pdf_path = pick_best_pdf(out_dir, year, kode)
        if pdf_path is None:
            logger.info("Tidak ada PDF terunduh untuk %s tahun %s, dilewati.", kode, year)
            continue
        try:
            profile = financials.extract_financial_profile(pdf_path)
        except Exception as exc:
            logger.warning("Gagal mengekstrak data keuangan %s (%s): %s", kode, pdf_path, exc)
            continue
        results.append((kode, periode, profile))
    return results


def _write_profil_keuangan_sheet(wb: Workbook, rows: list[tuple[str, str, financials.FinancialProfile]]) -> None:
    ws = wb.active
    ws.title = "Profil & Keuangan"
    headers = [
        "Kode", "Nama Entitas", "Periode Sumber", "Sektor", "Subsektor", "Industri", "Jenis Laporan",
        "Pendapatan (Rp juta)", "Pendapatan Periode Lalu (Rp juta)", "Pertumbuhan Pendapatan (%)",
        "Laba Bersih (Rp juta)", "Laba Bersih Periode Lalu (Rp juta)", "Pertumbuhan Laba (%)",
        "Beban Penjualan (Rp juta)", "Beban Umum & Administrasi (Rp juta)",
        "Kas & Setara Kas (Rp juta)", "Total Aset (Rp juta)",
        "Aset Lancar (Rp juta)", "Liabilitas Jangka Pendek (Rp juta)",
        "Jumlah Liabilitas (Rp juta)", "Jumlah Ekuitas (Rp juta)",
        "Current Ratio", "DER (Debt to Equity)", "Margin Laba Bersih (%)", "ROE (%)",
    ]
    _title_block(ws, "Profil & Keuangan (otomatis dari PDF laporan keuangan IDX)", (
        "Diekstrak otomatis oleh idx_report_scraper.intelligence dari PDF yang sudah diunduh. "
        "'Periode Sumber' menunjukkan periode mana yang datanya dipakai (diprioritaskan: tahunan > tw3 > tw2 > tw1). "
        "Current Ratio = Aset Lancar / Liabilitas Jangka Pendek (likuiditas jangka pendek, idealnya >= 1). "
        "DER = Jumlah Liabilitas / Jumlah Ekuitas (beban utang relatif terhadap modal sendiri, makin tinggi makin berisiko). "
        "Margin Laba Bersih = Laba Bersih / Pendapatan. ROE = Laba Bersih / Jumlah Ekuitas (efisiensi memutar modal). "
        "Keempatnya rumus Excel aktif, otomatis mengikuti kalau angka sumbernya diedit."
    ), len(headers))
    row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    first_data_row = row + 1
    row += 1
    for kode, periode, p in rows:
        values = [
            kode, p.entity_name, periode, p.sektor, p.subsektor, p.industri, p.jenis_laporan,
            p.pendapatan, p.pendapatan_prior, None,
            p.laba_bersih, p.laba_bersih_prior, None,
            p.beban_penjualan, p.beban_administrasi, p.kas_setara_kas, p.jumlah_aset,
            p.jumlah_aset_lancar, p.jumlah_liabilitas_jangka_pendek,
            p.jumlah_liabilitas, p.jumlah_ekuitas, None, None, None, None,
        ]
        for i, v in enumerate(values, start=1):
            ws.cell(row=row, column=i, value=v)
        ws.cell(row=row, column=10, value=f'=IFERROR((H{row}-I{row})/I{row},"-")')
        ws.cell(row=row, column=13, value=f'=IFERROR((K{row}-L{row})/L{row},"-")')
        # Current Ratio = Aset Lancar (R, col 18) / Liabilitas Jangka Pendek (S, col 19)
        ws.cell(row=row, column=22, value=f'=IFERROR(R{row}/S{row},"-")')
        # DER = Jumlah Liabilitas (T, col 20) / Jumlah Ekuitas (U, col 21)
        ws.cell(row=row, column=23, value=f'=IFERROR(T{row}/U{row},"-")')
        # Margin Laba Bersih = Laba Bersih (K, col 11) / Pendapatan (H, col 8)
        ws.cell(row=row, column=24, value=f'=IFERROR(K{row}/H{row},"-")')
        # ROE = Laba Bersih (K, col 11) / Jumlah Ekuitas (U, col 21)
        ws.cell(row=row, column=25, value=f'=IFERROR(K{row}/U{row},"-")')
        for c in (8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21):
            ws.cell(row=row, column=c).number_format = NUM_FMT
        for c in (10, 13, 24, 25):
            ws.cell(row=row, column=c).number_format = PCT_FMT
        for c in (22, 23):
            ws.cell(row=row, column=c).number_format = RATIO_FMT
        _style_data_row(ws, row, len(headers), wrap_cols=(2,))
        row += 1

    widths = [10, 32, 14, 22, 22, 24, 16, 18, 20, 16, 16, 20, 14, 18, 20, 18, 16, 16, 20, 18, 16, 12, 14, 16, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{first_data_row}"


def _write_matriks_kesehatan_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Matriks Kesehatan")
    headers = ["Revenue", "Profit", "Karyawan", "Cash", "Key", "Kategori", "Kondisi Kesehatan", "Arti untuk Commercial"]
    _title_block(ws, "Matriks Kesehatan Perusahaan - 16 Kombinasi", (
        "Revenue x Profit x Karyawan x Cash -> diagnosis kesehatan. B = Besar, K = Kecil, Byk = Banyak, Sdk = Sedikit "
        "(dibanding ambang di sheet Kalkulator Kesehatan)."
    ), len(headers))
    row = 4
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row, column=i, value=h)
    _style_header_row(ws, row, len(headers))
    row += 1
    for rev, prof, kar, cash, kategori, kondisi, arti in HEALTH_MATRIX:
        key = f"{rev}|{prof}|{kar}|{cash}"
        for i, v in enumerate([rev, prof, kar, cash, key, kategori, kondisi, arti], start=1):
            ws.cell(row=row, column=i, value=v)
        _style_data_row(ws, row, len(headers), wrap_cols=(7, 8))
        row += 1
    widths = [10, 10, 10, 10, 16, 12, 45, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"


def _write_kalkulator_sheet(wb: Workbook, rows: list[tuple[str, str, financials.FinancialProfile]]) -> None:
    ws = wb.create_sheet("Kalkulator Kesehatan")
    ncols = max(6, len(rows) + 1)
    _title_block(ws, "Kalkulator Kesehatan Perusahaan", (
        "Nilai dalam Rp miliar (dikonversi otomatis dari Rp juta hasil ekstraksi). 'Karyawan' TIDAK ada di laporan "
        "keuangan - isi manual (sel biru) kalau datanya tersedia dari sumber lain; kosongkan kalau tidak tahu."
    ), ncols)

    ws["A4"] = "AMBANG 'BESAR' (boleh diubah - Rp miliar / orang):"
    ws["A4"].font = Font(name=FONT_NAME, bold=True)
    thresholds = [("Revenue >=", 10000), ("Profit >=", 1000), ("Karyawan >=", 5000), ("Cash >=", 1000)]
    for i, (label, value) in enumerate(thresholds):
        r = 5 + i
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=value)
        cell.font = INPUT_FONT

    header_row = 10
    ws.cell(row=header_row, column=1, value="Item")
    for j, (kode, _periode, _p) in enumerate(rows, start=2):
        ws.cell(row=header_row, column=j, value=kode)
    _style_header_row(ws, header_row, 1 + len(rows))

    labels = [
        "Revenue (Rp miliar)", "Profit / laba bersih (Rp miliar)", "Karyawan (orang)", "Cash & setara kas (Rp miliar)",
        "Kelas Revenue", "Kelas Profit", "Kelas Karyawan", "Kelas Cash", "Key", "KATEGORI",
        "Kondisi kesehatan", "Arti untuk commercial",
    ]
    row_of = {name: header_row + 1 + i for i, name in enumerate(labels)}
    for name, r in row_of.items():
        ws.cell(row=r, column=1, value=name).font = Font(name=FONT_NAME, bold=(name in ("KATEGORI",)))

    for j, (kode, periode, p) in enumerate(rows, start=2):
        col = get_column_letter(j)
        rev_miliar = None if p.pendapatan is None else round(p.pendapatan / 1000, 3)
        profit_miliar = None if p.laba_bersih is None else round(p.laba_bersih / 1000, 3)
        cash_miliar = None if p.kas_setara_kas is None else round(p.kas_setara_kas / 1000, 3)

        ws.cell(row=row_of["Revenue (Rp miliar)"], column=j, value=rev_miliar)
        ws.cell(row=row_of["Profit / laba bersih (Rp miliar)"], column=j, value=profit_miliar)
        # Karyawan is intentionally left blank: not present in financial-statement PDFs.
        ws.cell(row=row_of["Karyawan (orang)"], column=j).font = INPUT_FONT
        ws.cell(row=row_of["Cash & setara kas (Rp miliar)"], column=j, value=cash_miliar)

        r = row_of["Kelas Revenue"]
        ws.cell(row=r, column=j, value=f'=IF({col}{row_of["Revenue (Rp miliar)"]}="","",IF({col}{row_of["Revenue (Rp miliar)"]}>=$B$5,"B","K"))')
        r = row_of["Kelas Profit"]
        ws.cell(row=r, column=j, value=f'=IF({col}{row_of["Profit / laba bersih (Rp miliar)"]}="","",IF({col}{row_of["Profit / laba bersih (Rp miliar)"]}>=$B$6,"B","K"))')
        r = row_of["Kelas Karyawan"]
        ws.cell(row=r, column=j, value=f'=IF({col}{row_of["Karyawan (orang)"]}="","",IF({col}{row_of["Karyawan (orang)"]}>=$B$7,"Byk","Sdk"))')
        r = row_of["Kelas Cash"]
        ws.cell(row=r, column=j, value=f'=IF({col}{row_of["Cash & setara kas (Rp miliar)"]}="","",IF({col}{row_of["Cash & setara kas (Rp miliar)"]}>=$B$8,"B","K"))')

        kr, kp, kk, kc = row_of["Kelas Revenue"], row_of["Kelas Profit"], row_of["Kelas Karyawan"], row_of["Kelas Cash"]
        r = row_of["Key"]
        ws.cell(row=r, column=j, value=(
            f'=IF(OR({col}{kr}="",{col}{kp}="",{col}{kk}="",{col}{kc}=""),"",'
            f'{col}{kr}&"|"&{col}{kp}&"|"&{col}{kk}&"|"&{col}{kc})'
        ))

        key_row = row_of["Key"]
        r = row_of["KATEGORI"]
        ws.cell(row=r, column=j, value=(
            f'=IFERROR(INDEX(\'Matriks Kesehatan\'!$F$5:$F$20,'
            f'MATCH({col}{key_row},\'Matriks Kesehatan\'!$E$5:$E$20,0)),"")'
        ))
        r = row_of["Kondisi kesehatan"]
        ws.cell(row=r, column=j, value=(
            f'=IFERROR(INDEX(\'Matriks Kesehatan\'!$G$5:$G$20,'
            f'MATCH({col}{key_row},\'Matriks Kesehatan\'!$E$5:$E$20,0)),"")'
        ))
        r = row_of["Arti untuk commercial"]
        ws.cell(row=r, column=j, value=(
            f'=IFERROR(INDEX(\'Matriks Kesehatan\'!$H$5:$H$20,'
            f'MATCH({col}{key_row},\'Matriks Kesehatan\'!$E$5:$E$20,0)),"")'
        ))

        for name in labels:
            cell = ws.cell(row=row_of[name], column=j)
            cell.border = BORDER
            if name not in ("Karyawan (orang)",):
                cell.font = NORMAL_FONT

    widths = [30] + [16] * len(rows)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B11"


def build_workbook(
    rows: list[tuple[str, str, financials.FinancialProfile]], out_path: Path
) -> None:
    """Build the Profil & Keuangan / Matriks Kesehatan / Kalkulator
    Kesehatan workbook from already-extracted (kode_emiten, periode,
    FinancialProfile) rows - see collect_financial_profiles()."""
    wb = Workbook()
    _write_profil_keuangan_sheet(wb, rows)
    _write_matriks_kesehatan_sheet(wb)
    _write_kalkulator_sheet(wb, rows)
    wb.save(out_path)
    logger.info("Workbook intelligence ditulis ke %s (%d perusahaan)", out_path, len(rows))
